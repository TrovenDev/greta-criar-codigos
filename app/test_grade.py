import sqlite3
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from grade_export import GRADE_HEADERS, grade_rows
from registration import Registration
from stock_io import export_xlsx


TEMPLATE = Path(__file__).resolve().parents[3] / 'dmaaa.xlsx'


class GradeTests(unittest.TestCase):
    def records(self):
        return [dict(codigo='27023112MCB',produto='27-0231 12 MCB',grupo='27-ANEL',referencia='AR38137-MRS04-F83 MCB 12',peso=16.54,unidade='G'),
                dict(codigo='27023114MO17',produto='27-0231 14 MO17',grupo='27-ANEL',referencia='AR38137-MRS04-F83 MO17 14',peso=1,unidade='G'),
                dict(codigo='26022004W',produto='26-0220 04 W',grupo='26-BRINCO',referencia='BR37647-MRS04-F83 W',peso=1,unidade='PR')]

    def test_unique_parents_at_the_end(self):
        rows = grade_rows(self.records())
        self.assertEqual(len(GRADE_HEADERS), 19)
        self.assertEqual(len(rows), 5)
        self.assertEqual([r[2] for r in rows[-2:]], ['270231','260220'])
        self.assertEqual([r[4] for r in rows[-2:]], ['AR38137-MRS04-F83','BR37647-MRS04-F83'])
        for r in rows[-2:]:
            self.assertEqual((r[7],r[8],r[13]),(0,50,0))
            self.assertEqual(r[11], '001')
            self.assertIsNone(r[17])
        self.assertEqual(rows[0][1], '0231')
        self.assertEqual(rows[0][6], '83')
        self.assertEqual(rows[0][13], 0.01654)
        self.assertEqual([r[9] for r in rows], ['G','G','PR','G','PR'])
        self.assertTrue(all(r[10]=='UN' for r in rows))

    def test_same_parent_cannot_refer_to_different_models(self):
        rows = self.records()
        rows[1]['referencia'] = 'AR-DIFFERENT-F83 W 14'
        with self.assertRaisesRegex(ValueError, 'referências-base diferentes'):
            grade_rows(rows)

    def test_parent_follows_children_for_all_units(self):
        for unit in ('PC', 'PR', 'G'):
            with self.subTest(unit=unit):
                records = [dict(r, unidade=unit) for r in self.records()[:2]]
                rows = grade_rows(records)
                self.assertEqual([r[9] for r in rows], [unit, unit, unit])
                self.assertEqual((rows[-1][7], rows[-1][8], rows[-1][13]), (0, 50, 0))

    def test_mixed_units_of_same_model_require_review(self):
        records = self.records()
        records[1]['unidade'] = 'PC'
        with self.assertRaisesRegex(ValueError, 'unidades diferentes'):
            grade_rows(records)

    def test_xlsx_columns_and_text_identifiers(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as tmp:
            path = Path(tmp) / 'Grade Lançar.xlsx'
            export_xlsx(path, self.records())
            b = load_workbook(path)
            s = b.active
            self.assertEqual(tuple(v.value for v in s[1]), GRADE_HEADERS)
            self.assertEqual(s.max_column,19)
            self.assertEqual(s['B2'].value,'0231')
            self.assertEqual(s['B2'].data_type,'s')
            self.assertEqual(s['S2'].value,'002')
            self.assertEqual(s['N2'].value,0.01654)
            self.assertEqual(s.max_row,6)
            b.close()

    def test_legacy_layout_matches_previous_grade_file(self):
        records = [dict(codigo='47095112W', produto='47-0951 12 W', grupo='47-ANEL MOISSANITERODIO', referencia='AR38850-MRS01-F23 W', peso=1, unidade='PC'),
                   dict(codigo='14095420Y102', produto='14-0954 20 Y102', grupo='14-ANEL', referencia='PLR37972-04-F83 BY102', peso=0.8, unidade='G')]
        registration = Registration(sqlite3.connect(':memory:'))
        registration.select_layout('legacy')
        headers, rows, name = registration.export(records)
        self.assertEqual(len(headers), 23)
        self.assertEqual(name, 'Grade Lançar · formato antigo')
        self.assertEqual(rows[0][:17], ['47095112W', '47-0951 12 W', 'PV', '01', 'PC', '47', '0951', '', '', 'A',
                                        '470951', '12', 'W', '71132000', '23', 'AR38850-MRS01-F23 W', '001'])
        self.assertEqual(rows[0][17:22], ['', 1, 1, 'UN', 0.001])
        self.assertEqual(rows[1][11:13], ['20', 'Y102'])
        self.assertEqual((rows[0][22], rows[1][22]), ('MOISSANITE', ''))
        self.assertEqual(len(rows), 2)

    def test_renamed_headers_keep_sources_and_order(self):
        registration = Registration(sqlite3.connect(':memory:'))
        registration.save_layout(dict(name='Minha grade', parents=False, columns=[
            dict(header='Modelo ERP', source='codigo'), dict(header='Fixo', source='fixed', value='001'), dict(header='Fator', source='erp.aaa_fatconv')]))
        records = self.records()[:2]
        headers, rows, _ = registration.export(records)
        self.assertEqual(headers, ['Modelo ERP', 'Fixo', 'Fator'])
        self.assertEqual(rows[0], ['27023112MCB', '001', 16.54])
        self.assertEqual(len(rows), 2)

    def test_default_layouts_can_be_renamed(self):
        registration = Registration(sqlite3.connect(':memory:'))
        columns = [dict(header='Código ERP', source='erp.aaa_codigo', value='')]
        registration.save_layout(dict(id='standard', name='Grade do ERP', parents=True, columns=columns))
        registration.select_layout('standard')
        headers, _, name = registration.export(self.records())
        self.assertEqual((headers, name), (['Código ERP'], 'Grade do ERP'))

    @unittest.skipUnless(TEMPLATE.exists(), 'Referência dmaaa.xlsx não fornecida neste ambiente')
    def test_matches_all_values_of_supplied_dmaaa(self):
        b = load_workbook(TEMPLATE,data_only=True)
        values = list(b.active.values)
        self.assertEqual(tuple(values[0]), GRADE_HEADERS)
        records = [dict(grupo=r[0],codigo=r[2],produto=r[3],referencia=r[4],peso=r[7],unidade=r[9]) for r in values[1:] if r[17]]
        actual = grade_rows(records)
        normalize = lambda row: [None if v=='' else v for v in row]
        self.assertEqual([normalize(r) for r in actual],[normalize(r) for r in values[1:]])
        b.close()
