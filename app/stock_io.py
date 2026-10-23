"""Leitura de estoque por cabeçalhos, independente do nome da aba."""
import re
import unicodedata
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


def header_key(text):
    text = unicodedata.normalize('NFKD', str(text or '').lower())
    return re.sub(r'[^a-z0-9]', '', ''.join(c for c in text if not unicodedata.combining(c)))


ALIASES = {
    'codigo': {'codsistema', 'codigosistema', 'codigodosistema', 'codigosistemico', 'aaacodigo', 'codigo', 'codproduto'},
    'produto': {'produto', 'descricao', 'descricaoproduto', 'aaadescricao', 'nomeproduto','identificadortecnico','produtotecnico'},
    'referencia': {'referencia', 'ref', 'referenciadofornecedor', 'codigofornecedor', 'aaacodrefforn', 'referenciafornecedor'},
    'grupo': {'gp', 'grupo', 'grupoproduto', 'aaagrupo','categoriagp'},
    'peso': {'peso', 'pesog', 'pesogramas', 'pesogr', 'aaapesogramas'},
    'refbase': {'aaarefbase','refbase','sequencia'},
}


def find_layout(sheet):
    for rownum, row in enumerate(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 30, 30), values_only=True), 1):
        keys = [header_key(v) for v in row]
        mapping = {name: next((i for i, key in enumerate(keys) if key in aliases), None) for name, aliases in ALIASES.items()}
        technical=next((i for i,key in enumerate(keys) if key in ('identificadortecnico','produtotecnico')),None)
        if technical is not None:mapping['produto']=technical
        if all(mapping[k] is not None for k in ('codigo', 'produto', 'referencia', 'grupo')):
            return rownum, mapping
    return None


def stock_sheets(path):
    with_book = load_workbook(path, read_only=True, data_only=True)
    try:
        return [s.title for s in with_book if find_layout(s)]
    finally:
        with_book.close()


def read_stock(path, sheet_name=None,known_records=None):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        available = [(s, find_layout(s)) for s in book if sheet_name is None or s.title == sheet_name]
        available = [(s, layout) for s, layout in available if layout]
        if not available:
            raise ValueError('Não encontrei os cabeçalhos de estoque: Cód. Sistema, Produto, Referência e GP. O nome da aba pode ser qualquer um.')
        if len(available) > 1:
            raise ValueError('Há mais de uma aba de estoque. Escolha a aba que deseja utilizar.')
        sheet, (head, mapping) = available[0]
        result = []
        known={str(r['codigo']).strip().upper():r for r in (known_records or [])}
        for n, row in enumerate(sheet.iter_rows(min_row=head + 1, values_only=True), head + 1):
            def get(key):
                i = mapping[key]
                return row[i] if i is not None and i < len(row) else None
            if get('codigo') in (None, ''):
                continue
            entry = {k: str(get(k) if get(k) is not None else '').strip() for k in ('codigo', 'produto', 'referencia', 'grupo')}
            if any(not entry[k] for k in ('produto', 'referencia', 'grupo')):
                raise ValueError(f'Aba {sheet.title}, linha {n}: preencha produto, referência e grupo antes de importar.')
            if not re.fullmatch(r'[A-Za-z0-9]+-\d+(?:\s+.*)?', entry['produto']):
                entry['descricao']=entry['produto']
                old=known.get(entry['codigo'].upper())
                if old and old['referencia'].strip().upper()==entry['referencia'].upper() and old['grupo'].split('-')[0]==entry['grupo'].split('-')[0]:
                    entry['produto']=old['produto']
                else:
                    prefix=entry['grupo'].split('-')[0].upper();raw=str(get('refbase') or '').strip()
                    code=entry['codigo'].upper()
                    sequence=next((s for s in (raw,raw.zfill(4)) if raw.isdigit() and code.startswith(prefix+s)),None)
                    if sequence:
                        tail=code[len(prefix+sequence):]
                        words=entry['referencia'].upper().split()
                        split_tail=(words[-1]+' '+words[1]) if len(words)==3 and words[-1].isdigit() else ''
                        if split_tail.replace(' ','')==tail:tail=split_tail
                        entry['produto']=prefix+'-'+sequence+(' '+tail if tail else '')
                    else:
                        raise ValueError(f"Aba {sheet.title}, linha {n}: não reconheço a sequência de {entry['produto']}. Inclua Identificador técnico (como 17-0523 14 W) ou aaa_refbase no relatório. Códigos reservados nesta base são reconhecidos automaticamente.")
            entry['peso'] = get('peso')
            result.append(entry)
        if not result:
            raise ValueError('A aba selecionada não possui produtos de estoque.')
        return result
    finally:
        book.close()


def export_xlsx(path, records, prepared=None,headers=None):
    """Exportador do aplicativo desktop; texto e números tipados, sem macros."""
    book = Workbook()
    sheet = book.active
    from grade_export import GRADE_HEADERS, grade_rows
    sheet.title = 'Planilha 1'
    sheet.append(list(headers if headers is not None else GRADE_HEADERS))
    rows = prepared if prepared is not None else grade_rows(records)
    for row in rows:
        sheet.append(row)
        for col, value in enumerate(row, 1):
            if isinstance(value, str):
                sheet.cell(sheet.max_row, col).data_type = 's'
                sheet.cell(sheet.max_row, col).number_format = '@'
    sheet.freeze_panes = 'A2'
    sheet.sheet_view.showGridLines = False
    for letter, width in [('A', 15), ('B', 16), ('C', 25), ('D', 26), ('E', 44), ('F', 13), ('G', 19), ('H', 19), ('I', 16), ('J', 12), ('K', 12), ('L', 21), ('M', 16), ('N', 15), ('O', 12), ('P', 15), ('Q', 14), ('R', 24), ('S', 19)]:
        sheet.column_dimensions[letter].width = width
    for row in sheet:
        sheet.row_dimensions[row[0].row].height = 24
        for cell in row:
            cell.font = Font(name='Arial', size=11)
            cell.alignment = Alignment(vertical='center', horizontal='right' if cell.column in (8, 9, 14) else 'left')
        if headers is None or list(headers)==list(GRADE_HEADERS):
            for index in (7, 8):row[index].number_format = '0.000###'
            row[13].number_format = '0.000000'
        if row[0].row > len(records) + 1:
            for cell in row:
                cell.fill = PatternFill('solid', fgColor='F7F5F1')
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='2B35F0')
        cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', horizontal='center')
    from openpyxl.utils import get_column_letter
    sheet.auto_filter.ref = f'A1:{get_column_letter(sheet.max_column)}{sheet.max_row}'
    book.save(path)
    book.close()
