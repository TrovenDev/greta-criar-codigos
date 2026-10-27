import unittest
from openpyxl import Workbook,load_workbook
import test_desktop as fixtures
from desktop_api import DesktopAPI
from stock_io import read_stock


class RegistrationTests(unittest.TestCase):
    setUp=fixtures.DesktopTests.setUp
    tearDown=fixtures.DesktopTests.tearDown
    ok=fixtures.DesktopTests.ok

    def seed(self):
        self.ok('manual',dict(text='NEW W 12\t\t\t47\nNEW W 14\t\t\t47'))

    def test_renamed_layout_exports_its_own_headers_and_fixed_values(self):
        self.seed()
        self.ok('layout_save',dict(name='Minha grade',parents=True,columns=[dict(header='Modelo',source='codigo'),dict(header='Constante',source='fixed',value='001'),dict(header='Fator',source='erp.aaa_fatconv')]))
        grade=self.ok('preview_grade')
        self.assertEqual(grade['headers'],['Modelo','Constante','Fator'])
        self.assertEqual(grade['rows'][-1],['471151','001',50])
        target=self.folder/'custom.xlsx';self.window.paths.append(target);self.ok('export')
        book=load_workbook(target);sheet=book.active
        self.assertEqual(sheet.max_column,3);self.assertEqual(sheet['B2'].value,'001');book.close()

    def test_custom_export_cannot_bypass_mixed_parent_units(self):
        self.seed();self.ok('layout_save',dict(name='Somente código',parents=False,columns=[dict(header='Código',source='codigo')]))
        self.ok('weight',dict(ids=[0],mode='Peso',value='2,35'))
        result=self.api.call('preview_grade');self.assertFalse(result['ok']);self.assertIn('unidades diferentes',result['error'])

    def test_written_stock_names_use_reserved_codes_or_explicit_sequence(self):
        source=self.folder/'nomes.xlsx';book=Workbook();sheet=book.active;sheet.title='Estoque'
        sheet.append(['Código','Produto','Referência','GP'])
        sheet.append(['47115112W','Anel solitário','NEW W 12','47-ANEL MOISSANITE']);book.save(source);book.close()
        with self.assertRaisesRegex(ValueError,'Identificador técnico'):read_stock(source)
        known=[dict(codigo='47115112W',produto='47-1151 12 W',referencia='NEW W 12',grupo='47-ANEL MOISSANITE')]
        self.assertEqual(read_stock(source,known_records=known)[0]['produto'],'47-1151 12 W')
        book=load_workbook(source);sheet=book.active;sheet.cell(1,5).value='aaa_refbase';sheet.cell(2,5).value='1151';book.save(source);book.close()
        self.assertEqual(read_stock(source)[0]['produto'],'47-1151 12 W')
