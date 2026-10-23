import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from desktop_api import DesktopAPI


class Dialogs:
    def __init__(self):
        self.paths=[]

    def create_file_dialog(self,*args,**kwargs):
        return [str(self.paths.pop(0))] if self.paths else None


class DesktopTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(dir=Path(__file__).parent)
        self.folder=Path(self.tmp.name)
        self.api=DesktopAPI(self.folder)
        self.window=Dialogs()
        self.api._window=self.window
        self.mock=patch.dict(sys.modules,webview=types.SimpleNamespace(FileDialog=types.SimpleNamespace(OPEN=1,SAVE=2)))
        self.mock.start()
        self.stock=self.folder/'estoque.xlsx'
        b=Workbook();s=b.active;s.title='Estoque'
        s.append(['Cód. Sistema','Produto','Referência','GP'])
        s.append(['27023014MO5','27-0230 14 MO5','AR38138-MRS03-F83 MO5 14','27-ANEL PRATA'])
        s.append(['47023014W','47-0230 14 W','AR38138-MRS03-F83 W 14','47-ANEL MOISSANITE'])
        s.append(['47115014W','47-1150 14 W','OLD W 14','47-ANEL MOISSANITE'])
        b.save(self.stock);b.close()
        self.window.paths.append(self.stock)
        self.ok('choose_stock')
        self.ok('load_stock',dict(sheet='Estoque'))

    def tearDown(self):
        self.api._close();self.mock.stop();self.tmp.cleanup()

    def ok(self,action,p=None):
        result=self.api.call(action,p)
        self.assertTrue(result['ok'],result)
        return result['data']

    def test_category_application_marks_former_pending_for_export(self):
        self.ok('manual',dict(text='NEW W 12\nNEW W 14'))
        d=self.ok('category',dict(ids=[0],group='47-ANEL MOISSANITE',model=True))
        self.assertEqual([r['produto'] for r in d['rows']],['47-1151 12 W','47-1151 14 W'])
        self.assertEqual([r['situacao'] for r in d['rows']],['Novo modelo','Nova variação'])
        self.assertEqual(d['counts']['exportar'],2)
        self.ok('toggle',dict(ids=[1],value=False))
        d=self.ok('category',dict(ids=[0],group='47-ANEL MOISSANITE'))
        self.assertEqual(d['counts']['exportar'],1)

    def test_moissanite_cannot_be_assigned_prata(self):
        d=self.ok('manual',dict(text='AR38138-MRS03-F83 MO5 14'))
        self.assertEqual(d['rows'][0]['codigo'],'47023014MO5')
        r=self.api.call('category',dict(ids=[0],group='27-ANEL PRATA'))
        self.assertFalse(r['ok']);self.assertIn('moissanite',r['error'])
        self.assertEqual(self.ok('state')['rows'][0]['grupo'],'47-ANEL MOISSANITE')

    def test_import_export_parent_unit_and_restore(self):
        incoming=self.folder/'novos.xlsx';b=Workbook();s=b.active;s.title='Novos'
        s.append(['Referência','Peso','GP','Não usar'])
        s.append(['NEW W 12',None,'47','irrelevante']);s.append(['NEW W 14',None,'47','irrelevante'])
        b.save(incoming);b.close();self.window.paths.append(incoming)
        info=self.ok('choose_new');self.assertEqual(info['inspection']['mapping']['referencia'],0)
        d=self.ok('import_new',dict(sheet='Novos',header=1,mapping=info['inspection']['mapping']))
        self.assertEqual(d['imported'],2);self.assertEqual(d['state']['counts']['exportar'],2)
        grade=self.ok('preview_grade');self.assertEqual((grade['children'],grade['parents']),(2,1))
        self.assertEqual(grade['rows'][-1][9],'PC')
        target=self.folder/'Grade Lançar.xlsx';self.window.paths.append(target)
        self.ok('export');b=load_workbook(target);self.assertEqual(b.active['J4'].value,'PC');self.assertEqual(b.active['I4'].value,50);b.close()
        d=self.ok('restore');self.assertTrue(all(r['status']=='Reservado' for r in d['rows']))
        self.assertTrue(all(r['unidade']=='PC' for r in d['rows']))

    def test_mixed_units_block_before_reservation(self):
        self.ok('manual',dict(text='NEW W 12\t\t\t47\nNEW W 14\t\t\t47'))
        self.ok('weight',dict(ids=[0],mode='Peso',value='2,35'))
        result=self.api.call('export')
        self.assertFalse(result['ok']);self.assertIn('unidades diferentes',result['error'])
        count=self.api._pool.submit(lambda:self.api._catalog.db.execute('SELECT COUNT(*) FROM reservations').fetchone()[0]).result()
        self.assertEqual(count,0)

    def test_rule_applies_to_batch_and_persists(self):
        self.ok('manual',dict(text='AR-NEW W 12\nAR-NEW W 14'))
        d=self.ok('rule_save',dict(pattern='AR-NEW*',group='47-ANEL MOISSANITE'))
        self.assertEqual(d['counts']['exportar'],2)
        self.assertEqual(len(d['rules']),1)
        d=self.ok('restore');self.assertEqual(len(d['rules']),1)

    def test_cancel_keeps_current_state(self):
        self.ok('manual',dict(text='NEW W 12\t\t\t47'))
        self.assertIsNone(self.ok('choose_new'))
        self.assertEqual(self.ok('state')['counts']['total'],1)
        self.assertFalse(self.api.call('invalid_action')['ok'])
