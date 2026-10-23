"""Importação de itens novos com mapeamento explícito e prévia."""
import re
from openpyxl import load_workbook
from stock_io import header_key
from item_rules import weight, piece_unit

FIELDS = ('referencia', 'peso', 'grupo', 'tipo', 'tamanho_cor')
LABELS = {'referencia': 'Referência (obrigatório)', 'peso': 'Peso (opcional)', 'grupo': 'Categoria / GP (opcional)', 'tipo': 'Tipo de produto (opcional)', 'tamanho_cor': 'Tamanho e cor (opcional)'}
ALIASES = {
    'referencia': ('referencia', 'ref', 'referenciadofornecedor', 'aaacodrefforn', 'codigofornecedor', 'codigoreferencia'),
    'peso': ('peso', 'pesog', 'pesogramas', 'aaapesogramas'),
    'grupo': ('gp', 'grupo', 'categoria', 'aaagrupo'),
    'tipo': ('tipo', 'tipoproduto', 'tipodoproduto', 'categoriaitem'),
    'tamanho_cor': ('tamanhoecor', 'sufixo', 'variacao'),
}


def list_sheets(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def inspect_sheet(path, sheet_name, header_row=None):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[sheet_name]
        top = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 30, 30), values_only=True))
        if header_row is None:
            def score(row):
                keys = {header_key(v) for v in row}
                return 10 * bool(keys.intersection(ALIASES['referencia'])) + sum(bool(keys.intersection(ALIASES[f])) for f in FIELDS[1:])
            header_row = max(range(len(top)), key=lambda i: score(top[i])) + 1 if top else 1
        values = next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
        columns = [f'{i + 1:02d} — {v if v is not None else "Coluna sem título"}' for i, v in enumerate(values)]
        keys = [header_key(v) for v in values]
        mapping = {field: next((i for alias in ALIASES[field] for i, key in enumerate(keys) if key == alias), None) for field in FIELDS}
        if mapping['referencia'] is None:
            mapping['referencia'] = next((i for i, key in enumerate(keys) if key in ('codigo', 'produto', 'descricao', 'nome')), None)
        preview = list(sheet.iter_rows(min_row=header_row + 1, max_row=min(header_row + 8, sheet.max_row or header_row + 8), values_only=True))
        return header_row, columns, mapping, preview
    finally:
        book.close()


def read_items(path, sheet_name, header_row, mapping):
    if mapping.get('referencia') is None:
        raise ValueError('Selecione a coluna que contém a referência dos produtos.')
    used = [v for v in mapping.values() if v is not None]
    if len(used) != len(set(used)):
        raise ValueError('Uma mesma coluna não pode representar dois campos. Confira o mapeamento.')
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        result, skipped = [], 0
        for line, cells in enumerate(book[sheet_name].iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            def get(field):
                index = mapping.get(field)
                return cells[index] if index is not None and index < len(cells) else None
            ref = str(get('referencia') or '').strip()
            if not ref or header_key(ref) in ('total', 'totais', 'subtotal', 'totalgeral', 'valortotal'):
                skipped += 1
                continue
            raw_weight = get('peso')
            try:
                numeric = weight(raw_weight)
            except ValueError as exc:
                raise ValueError(f'Linha {line}, referência {ref}: {exc}') from None
            record = {field: str(get(field) or '').strip() for field in FIELDS if field != 'peso'}
            record.update(peso=numeric, linha_origem=line)
            result.append(record)
        if not result:
            raise ValueError('Não encontrei produtos na coluna de referência escolhida.')
        return result, skipped
    finally:
        book.close()


class ImportDialog:
    def __init__(self, app, path, accept):
        import tkinter as tk
        from tkinter import ttk
        import troven_brand as brand
        self.app, self.path, self.accept = app, path, accept
        self.window = tk.Toplevel(app.root)
        self.window.title('Troven • Importar planilha de códigos novos')
        screen_w, screen_h = app.root.winfo_screenwidth(), app.root.winfo_screenheight()
        width = min(1120, screen_w - 60)
        height = min(760, screen_h - 140)
        self.window.geometry(f'{width}x{height}+{max(0, (screen_w-width)//2)}+{max(0, (screen_h-height)//2-25)}')
        self.window.minsize(min(620, width), min(380, height))
        self.window.transient(app.root)
        self.window.configure(bg=brand.BRANCO)
        self.window.rowconfigure(0, weight=1)
        self.window.columnconfigure(0, weight=1)
        # O rodapé tem seu próprio espaço; o conteúdo pode rolar sem esconder a ação.
        self.footer = ttk.Frame(self.window, padding=(16, 12))
        self.footer.grid(row=1, column=0, sticky='ew')
        self.import_button = ttk.Button(self.footer, text='Importar produtos', style='Primary.TButton', command=self.confirm)
        self.import_button.pack(side='right')
        self.cancel_button = ttk.Button(self.footer, text='Cancelar', command=self.window.destroy)
        self.cancel_button.pack(side='right', padx=10)
        body = ttk.Frame(self.window)
        body.grid(row=0, column=0, sticky='nsew')
        body.rowconfigure(0, weight=1)
        body.columnconfigure(0, weight=1)
        self.canvas = tk.Canvas(body, highlightthickness=0, bg=brand.BRANCO)
        self.canvas.grid(row=0, column=0, sticky='nsew')
        vertical = ttk.Scrollbar(body, orient='vertical', command=self.canvas.yview)
        vertical.grid(row=0, column=1, sticky='ns')
        horizontal = ttk.Scrollbar(body, orient='horizontal', command=self.canvas.xview)
        horizontal.grid(row=1, column=0, sticky='ew')
        self.canvas.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        page = ttk.Frame(self.canvas, padding=16)
        content = self.canvas.create_window(0, 0, window=page, anchor='nw')

        def resize_content(event=None):
            self.canvas.itemconfigure(content, width=max(self.canvas.winfo_width(), page.winfo_reqwidth()),
                                      height=max(self.canvas.winfo_height(), page.winfo_reqheight()))
            self.canvas.configure(scrollregion=self.canvas.bbox('all'))

        page.bind('<Configure>', resize_content)
        self.canvas.bind('<Configure>', resize_content)

        def scroll(event):
            if not isinstance(event.widget, (ttk.Treeview, ttk.Combobox, ttk.Spinbox)):
                self.canvas.yview_scroll(-int(event.delta / 120), 'units')
                return 'break'

        self.window.bind('<MouseWheel>', scroll)
        self.window.bind('<Escape>', lambda e: self.window.destroy())
        self.window.bind('<Alt-i>', lambda e: self.confirm())
        ttk.Label(page, text='Confira o que será importado.', style='PageTitle.TLabel').pack(anchor='w')
        ttk.Label(page, text='Escolha as colunas úteis. As demais não entram no cadastro.', style='Muted.TLabel').pack(anchor='w', pady=(8, 12))
        bar = ttk.Frame(page)
        bar.pack(fill='x')
        ttk.Label(bar, text='Aba:').pack(side='left')
        self.sheet = tk.StringVar()
        combo = ttk.Combobox(bar, textvariable=self.sheet, values=list_sheets(path), state='readonly', width=30)
        combo.pack(side='left', padx=(6, 16))
        self.sheet.set(combo['values'][0])
        combo.bind('<<ComboboxSelected>>', lambda e: self.inspect())
        ttk.Label(bar, text='Linha dos títulos:').pack(side='left')
        self.header = tk.IntVar(value=1)
        ttk.Spinbox(bar, from_=1, to=1000, textvariable=self.header, width=6).pack(side='left', padx=6)
        ttk.Button(bar, text='Ler títulos', command=lambda: self.inspect(self.header.get())).pack(side='left')
        selectors = ttk.Frame(page)
        selectors.pack(fill='x', pady=14)
        self.vars, self.selectors = {}, {}
        for i, field in enumerate(FIELDS):
            r, c = divmod(i, 3)
            cell = ttk.Frame(selectors)
            cell.grid(row=r, column=c, sticky='ew', padx=(0, 12), pady=5)
            ttk.Label(cell, text=LABELS[field]).pack(anchor='w')
            self.vars[field] = tk.StringVar()
            pick = ttk.Combobox(cell, textvariable=self.vars[field], state='readonly', width=29)
            pick.pack(fill='x', pady=(5, 0))
            self.selectors[field] = pick
            pick.bind('<<ComboboxSelected>>', lambda e: self.preview())
            selectors.columnconfigure(c, weight=1)
        ttk.Label(page, text='Peso em branco ou coluna de peso não selecionada: valor 1 e venda por PC/PR. Peso informado: unidade G.', style='Muted.TLabel', wraplength=1020).pack(anchor='w')
        self.tree = app.make_tree(page, ('referencia', 'peso', 'unidade', 'grupo', 'tipo'), ('Referência', 'Peso', 'Unidade', 'Categoria / GP da origem', 'Tipo da origem'), (300, 100, 90, 250, 190))
        self.tree.configure(height=5)
        self.info = tk.StringVar()
        ttk.Label(page, textvariable=self.info, style='Muted.TLabel', wraplength=1020).pack(anchor='w', pady=10)
        self.inspect()

    def inspect(self, header=None):
        try:
            row, self.columns, mapping, self.raw_preview = inspect_sheet(self.path, self.sheet.get(), header)
            self.header.set(row)
            choices = ['Não utilizar'] + self.columns
            for field in FIELDS:
                self.selectors[field]['values'] = choices
                self.vars[field].set(self.columns[mapping[field]] if mapping[field] is not None else 'Não utilizar')
            self.preview()
        except Exception as exc:
            self.app.message('Não foi possível ler a aba', exc, True)

    def mapping(self):
        return {f: self.columns.index(v.get()) if v.get() in self.columns else None for f, v in self.vars.items()}

    def preview(self):
        mapping = self.mapping()
        self.tree.delete(*self.tree.get_children())
        count = 0
        for row in self.raw_preview:
            def get(field):
                index = mapping[field]
                return row[index] if index is not None and index < len(row) else None
            ref = str(get('referencia') or '').strip()
            if not ref or header_key(ref) in ('total', 'totais', 'subtotal', 'totalgeral', 'valortotal'):
                continue
            try:
                number = weight(get('peso'))
                unit = 'G' if number is not None else piece_unit(ref, str(get('grupo') or ''), str(get('tipo') or ''))
                value = number if number is not None else 1
            except ValueError:
                value, unit = 'Peso inválido', 'Revisar'
            self.tree.insert('', 'end', values=(ref, value, unit, get('grupo') or '', get('tipo') or ''))
            count += 1
        self.info.set(f'Prévia das primeiras {count} linhas. Confira a referência e o peso antes de importar. Linhas sem referência e totais identificados serão ignorados.')

    def confirm(self):
        try:
            records, skipped = read_items(self.path, self.sheet.get(), self.header.get(), self.mapping())
            self.accept(records, self.path, self.sheet.get(), skipped)
            self.window.destroy()
        except Exception as exc:
            self.app.message('Planilha não importada', exc, True)
