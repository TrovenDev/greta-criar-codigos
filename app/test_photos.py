import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from photo_links import PhotoLinks
from grade_export import grade_rows
from stock_io import export_xlsx


class PhotoTests(unittest.TestCase):
    def test_exact_codes_parent_and_missing_export(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as tmp:
            folder=Path(tmp);photos=folder/'images';photos.mkdir()
            for filename in ('97089318TOP.jpg','97089318TOP.png','970893.jpg','unrelated.jpg'):
                (photos/filename).write_bytes(b'photo')
            rows=[dict(codigo='97089318TOP',produto='97-0893 18 TOP',referencia='AR-A TOP 18',peso=1,unidade='PC'),
                  dict(codigo='97089320TOP',produto='97-0893 20 TOP',referencia='AR-A TOP 20',peso=1,unidade='PC')]
            links=PhotoLinks(folder);links.configure(photos);summary=links.apply(rows,strict=True,force=True)
            self.assertEqual(summary['found'],1);self.assertEqual(summary['missing'],1)
            grade=grade_rows(rows)
            self.assertEqual(grade[0][15],str(photos/'97089318TOP.jpg'))
            self.assertFalse(grade[1][15]);self.assertEqual(grade[-1][15],str(photos/'970893.jpg'))
            path=folder/'Grade.xlsx';export_xlsx(path,rows,prepared=grade)
            wb=load_workbook(path);self.assertEqual(wb.active.cell(2,16).value,str(photos/'97089318TOP.jpg'));wb.close()
            (photos/'97089318TOP.jpg').unlink();links.apply(rows,force=True)
            self.assertTrue(rows[0]['foto'].endswith('.png'))
            (photos/'97089318TOP.png').unlink();links.apply(rows,force=True)
            self.assertFalse(rows[0]['foto'])
            links.configure(photos,False);links.apply(rows)
            self.assertFalse(rows[0]['foto_pai'])

    def test_unavailable_enabled_folder_blocks_export(self):
        with tempfile.TemporaryDirectory(dir=Path(__file__).parent) as tmp:
            links=PhotoLinks(tmp);links.enabled=True;links.folder=str(Path(tmp)/'missing')
            with self.assertRaisesRegex(ValueError,'pasta de fotos'):links.apply([],strict=True)
